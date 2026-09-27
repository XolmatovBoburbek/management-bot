"""Чтение ERP-таблицы (xlsx / выгрузка Google Sheets) и запись статусов обратно.

Таблица «сырая» и будет меняться, поэтому колонки ищутся по заголовкам,
а не по буквам. Листы определяются по содержимому, а не по названию.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from app.dates import parse_date
from app.models import (
    STATUS_LABELS,
    Task,
    norm_text,
    normalize_priority,
    normalize_status,
)

# Поле -> варианты заголовка. Порядок важен: сначала точное совпадение,
# потом «заголовок начинается с».
TASK_COLUMNS: dict[str, list[str]] = {
    "no": ["№", "no", "#", "номер", "n"],
    "block": ["блок", "раздел", "направление", "этап"],
    "title": ["процедура", "задача", "название задачи", "наименование"],
    "description": ["что конкретно сделать", "описание", "что сделать", "детали"],
    "priority": ["приоритет", "важность"],
    "responsible": ["ответственный", "ответственные", "ответственное лицо"],
    "status": ["статус процедуры", "статус задачи", "статус"],
    "start_date": ["начало работы", "дата начала", "начало", "старт"],
    "deadline": ["окончание работы", "дедлайн", "дата окончания", "срок выполнения", "окончание", "deadline"],
    "fact_date": ["фактическая дата завершения", "фактическая дата", "факт"],
    "contractor": ["поставщик / исполнитель", "поставщик/исполнитель", "поставщик", "подрядчик"],
    "proof": ["доказательство выполнения", "доказательство", "результат"],
    "comment": ["комментарии исполнителя", "комментарий исполнителя", "комментарии", "комментарий"],
}
DATE_FIELDS = ("start_date", "deadline", "fact_date")
MAX_EMPTY_ROWS = 30


@dataclass
class ParsedTask:
    row_index: int
    title: str
    no: str = ""
    block: str = ""
    description: str = ""
    priority: str = ""
    responsible: str = ""
    status: str = "todo"
    status_raw: str = ""
    start_date: date | None = None
    deadline: date | None = None
    fact_date: date | None = None
    contractor: str = ""
    proof: str = ""
    comment: str = ""


@dataclass
class ParsedMilestone:
    title: str
    date_raw: str = ""
    criteria: str = ""
    status: str = ""


@dataclass
class ParsedRisk:
    title: str
    no: str = ""
    probability: str = ""
    impact: str = ""
    scenario: str = ""
    mitigation: str = ""
    owner: str = ""
    status: str = ""


@dataclass
class ParsedWorkbook:
    code: str
    name: str
    event_date: date | None
    info: list[list[str]] = field(default_factory=list)
    tasks: list[ParsedTask] = field(default_factory=list)
    milestones: list[ParsedMilestone] = field(default_factory=list)
    risks: list[ParsedRisk] = field(default_factory=list)
    layout: dict = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


class WorkbookError(ValueError):
    pass


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return "" if text.startswith("=") else text


def _match_header(header: str, field_name: str, exact: bool) -> bool:
    for option in TASK_COLUMNS[field_name]:
        if exact and header == option:
            return True
        if not exact and len(option) > 3 and header.startswith(option):
            return True
    return False


def detect_task_columns(ws: Worksheet, max_scan_rows: int = 15) -> tuple[int, dict[str, int]] | None:
    """Находит строку заголовков задач и сопоставление поле -> номер колонки (1-based)."""
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, max_scan_rows)):
        headers = [(cell.column, norm_text(cell.value)) for cell in row if cell.value is not None]
        if not headers:
            continue
        mapping: dict[str, int] = {}
        for exact in (True, False):
            for field_name in TASK_COLUMNS:
                if field_name in mapping:
                    continue
                for col, header in headers:
                    if col in mapping.values():
                        continue
                    if _match_header(header, field_name, exact):
                        mapping[field_name] = col
                        break
        if "title" in mapping and "responsible" in mapping:
            return row[0].row, mapping
    return None


def _parse_info(ws: Worksheet) -> list[list[str]]:
    """Пары «Метка: значение» с листа «Информация о проекте»."""
    pairs: list[list[str]] = []
    for row in ws.iter_rows():
        cells = [(c.column, c.value) for c in row if c.value not in (None, "")]
        if len(cells) < 2:
            continue
        label = _cell_text(cells[0][1])
        if not label.endswith(":"):
            continue
        value = cells[1][1]
        if hasattr(value, "strftime"):
            value = value.strftime("%d.%m.%Y")
        pairs.append([label.rstrip(":").strip(), _cell_text(value)])
    return pairs


def _info_value(info: list[list[str]], *keys: str) -> str:
    for label, value in info:
        text = norm_text(label)
        if any(key in text for key in keys):
            return value
    return ""


def _find_header_row(ws: Worksheet, required: list[str], max_scan_rows: int = 10) -> tuple[int, dict[str, int]] | None:
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, max_scan_rows)):
        headers = {norm_text(c.value): c.column for c in row if c.value is not None}
        if all(any(req in h for h in headers) for req in required):
            return row[0].row, headers
    return None


def _col(headers: dict[str, int], *keys: str) -> int | None:
    for key in keys:
        for header, col in headers.items():
            if header.startswith(key):
                return col
    return None


def _parse_milestones(ws: Worksheet) -> list[ParsedMilestone] | None:
    found = _find_header_row(ws, ["milestone", "дата"])
    if not found:
        found = _find_header_row(ws, ["веха", "дата"])
    if not found:
        return None
    header_row, headers = found
    c_title = _col(headers, "milestone", "веха")
    c_date = _col(headers, "дата")
    c_crit = _col(headers, "критери")
    c_status = _col(headers, "статус")
    items = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        get = lambda c: _cell_text(row[c - 1]) if c and c - 1 < len(row) else ""  # noqa: E731
        title = get(c_title)
        if not title:
            continue
        date_value = row[c_date - 1] if c_date and c_date - 1 < len(row) else None
        date_raw = date_value.strftime("%d.%m.%Y") if hasattr(date_value, "strftime") else _cell_text(date_value)
        items.append(ParsedMilestone(title=title, date_raw=date_raw, criteria=get(c_crit), status=get(c_status)))
    return items


def _parse_risks(ws: Worksheet) -> list[ParsedRisk] | None:
    found = _find_header_row(ws, ["риск", "ответствен"])
    if not found:
        return None
    header_row, headers = found
    cols = {
        "no": _col(headers, "№"),
        "title": _col(headers, "риск"),
        "probability": _col(headers, "вероятн"),
        "impact": _col(headers, "влияни"),
        "scenario": _col(headers, "сценари"),
        "mitigation": _col(headers, "решени", "митигац", "что делать"),
        "owner": _col(headers, "ответствен"),
        "status": _col(headers, "статус"),
    }
    items = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        values = {k: (_cell_text(row[c - 1]) if c and c - 1 < len(row) else "") for k, c in cols.items()}
        if not values["title"]:
            continue
        items.append(ParsedRisk(**values))
    return items


def _parse_tasks(ws: Worksheet, header_row: int, mapping: dict[str, int], event_date: date | None,
                 warnings: list[str]) -> list[ParsedTask]:
    tasks: list[ParsedTask] = []
    empty_streak = 0
    for row in ws.iter_rows(min_row=header_row + 1):
        by_col = {cell.column: cell.value for cell in row}
        title = _cell_text(by_col.get(mapping["title"]))
        if not title:
            empty_streak += 1
            if empty_streak >= MAX_EMPTY_ROWS:
                break
            continue
        empty_streak = 0
        row_index = row[0].row

        def text(field_name: str) -> str:
            col = mapping.get(field_name)
            return _cell_text(by_col.get(col)) if col else ""

        dates: dict[str, date | None] = {}
        for field_name in DATE_FIELDS:
            col = mapping.get(field_name)
            raw = by_col.get(col) if col else None
            parsed = parse_date(raw, event_date)
            if raw not in (None, "") and parsed is None and norm_text(raw) not in {"уточнить", "-", "—", "tbd"}:
                if not (isinstance(raw, str) and raw.startswith("=")):
                    warnings.append(f"Строка {row_index}: не удалось распознать дату «{raw}»")
            dates[field_name] = parsed

        status_raw = text("status")
        tasks.append(
            ParsedTask(
                row_index=row_index,
                title=title,
                no=text("no"),
                block=text("block"),
                description=text("description"),
                priority=normalize_priority(text("priority")),
                responsible=text("responsible"),
                status=normalize_status(status_raw),
                status_raw=status_raw,
                contractor=text("contractor"),
                proof=text("proof"),
                comment=text("comment"),
                **dates,
            )
        )
    return tasks


def parse_workbook(source: str | Path | bytes, fallback_name: str = "Проект") -> ParsedWorkbook:
    try:
        data = io.BytesIO(source) if isinstance(source, bytes) else source
        wb = openpyxl.load_workbook(data, data_only=True)
    except Exception as exc:  # noqa: BLE001 — openpyxl бросает разные исключения
        raise WorkbookError(f"Не удалось открыть файл как Excel (.xlsx): {exc}") from exc

    warnings: list[str] = []
    info: list[list[str]] = []
    tasks_ws: tuple[Worksheet, int, dict[str, int]] | None = None
    milestones: list[ParsedMilestone] = []
    risks: list[ParsedRisk] = []
    layout: dict = {}

    for ws in wb.worksheets:
        if tasks_ws is None:
            detected = detect_task_columns(ws)
            if detected:
                tasks_ws = (ws, detected[0], detected[1])
                continue
        name = norm_text(ws.title)
        if not info and ("информац" in name or "о проекте" in name):
            info = _parse_info(ws)
            continue
        if not milestones:
            parsed_ms = _parse_milestones(ws)
            if parsed_ms is not None:
                milestones = parsed_ms
                layout["milestones_sheet"] = ws.title
                continue
        if not risks:
            parsed_risks = _parse_risks(ws)
            if parsed_risks is not None:
                risks = parsed_risks
                layout["risks_sheet"] = ws.title
                continue
        if not info:
            candidate = _parse_info(ws)
            if _info_value(candidate, "дата мероприят", "код проекта"):
                info = candidate

    if tasks_ws is None:
        raise WorkbookError(
            "Не нашёл лист с задачами. Нужны колонки хотя бы «Процедура» (или «Задача») и «Ответственный»."
        )

    event_date = parse_date(_info_value(info, "дата мероприят"))
    ws, header_row, mapping = tasks_ws
    tasks = _parse_tasks(ws, header_row, mapping, event_date, warnings)
    layout.update({"tasks_sheet": ws.title, "header_row": header_row, "columns": mapping})

    code = _info_value(info, "код проекта") or re.sub(r"[^\w]+", "_", fallback_name).strip("_").upper()
    name = _info_value(info, "название мероприят", "название проекта") or fallback_name
    return ParsedWorkbook(
        code=code or "PROJECT",
        name=name,
        event_date=event_date,
        info=info,
        tasks=tasks,
        milestones=milestones,
        risks=risks,
        layout=layout,
        warnings=warnings,
    )


def comment_lines_to_append(existing: str, lines: list[str]) -> list[str]:
    """Строки комментариев из бота, которых ещё нет в ячейке таблицы."""
    existing_norm = norm_text(existing)
    return [line for line in lines if norm_text(line) not in existing_norm]


STANDARD_HEADERS = [
    ("no", "№", 5),
    ("block", "Блок", 22),
    ("title", "Процедура", 38),
    ("description", "Что конкретно сделать", 50),
    ("priority", "Приоритет", 12),
    ("responsible", "Ответственный", 16),
    ("status", "Статус процедуры", 22),
    ("start_date", "Начало работы", 14),
    ("deadline", "Окончание работы", 14),
    ("fact_date", "Фактическая дата завершения", 16),
    ("contractor", "Поставщик / Исполнитель", 28),
    ("proof", "Доказательство выполнения", 28),
    ("comment", "Комментарии исполнителя", 40),
]


def _write_task_row(ws: Worksheet, row: int, cols: dict[str, int], task: Task, today: date,
                    comment_lines: list[str], full: bool) -> None:
    """full=True — строка новая: пишем все поля; иначе только то, что меняется в боте."""
    def put(field_name: str, value: object) -> None:
        if field_name in cols:
            ws.cell(row, cols[field_name]).value = value

    if full:
        put("no", int(task.no) if task.no.isdigit() else task.no)
        put("block", task.block)
        put("title", task.title)
        put("description", task.description)
        put("contractor", task.contractor)
        put("proof", task.proof)
    put("status", STATUS_LABELS[task.effective_status(today)])
    for field_name in DATE_FIELDS:
        if field_name in cols:
            value = getattr(task, field_name)
            cell = ws.cell(row, cols[field_name])
            cell.value = value
            if value:
                cell.number_format = "DD.MM.YYYY"
    if task.responsible or full:
        put("responsible", task.responsible)
    if task.priority or full:
        put("priority", task.priority)
    if "comment" in cols:
        cell = ws.cell(row, cols["comment"])
        current = task.sheet_comment if full else _cell_text(cell.value)
        extra = comment_lines_to_append(current, comment_lines)
        if extra or full:
            cell.value = "\n".join([current, *extra]).strip() or None


def export_workbook(original: str | Path, layout: dict, tasks: list[Task], today: date,
                    bot_comments: dict[int, list[str]] | None = None) -> bytes:
    """Записывает статусы/даты/комментарии из бота в исходный файл и возвращает xlsx.

    Задачи, созданные в боте (их нет в файле), дописываются новыми строками в конец списка.
    """
    from copy import copy

    wb = openpyxl.load_workbook(original)
    sheet_name = layout.get("tasks_sheet")
    if sheet_name not in wb.sheetnames:
        raise WorkbookError("В исходном файле нет листа с задачами — загрузите таблицу заново.")
    ws = wb[sheet_name]
    cols: dict[str, int] = {k: int(v) for k, v in layout.get("columns", {}).items()}
    header_row = int(layout.get("header_row", 1))
    title_col = cols["title"]

    rows_by_title: dict[str, int] = {}
    last_row = header_row
    max_no = 0
    for r in range(header_row + 1, ws.max_row + 1):
        title = norm_text(ws.cell(r, title_col).value)
        if title:
            rows_by_title.setdefault(title, r)
            last_row = r
            if "no" in cols and str(ws.cell(r, cols["no"]).value or "").isdigit():
                max_no = max(max_no, int(ws.cell(r, cols["no"]).value))

    bot_comments = bot_comments or {}
    template_row = last_row if last_row > header_row else None
    for task in tasks:
        row = task.row_index
        if not row or norm_text(ws.cell(row, title_col).value) != norm_text(task.title):
            row = rows_by_title.get(norm_text(task.title))
        if row:
            _write_task_row(ws, row, cols, task, today, bot_comments.get(task.id, []), full=False)
            continue
        last_row += 1
        if template_row:
            for col in cols.values():
                src = ws.cell(template_row, col)
                dst = ws.cell(last_row, col)
                if src.has_style:
                    dst._style = copy(src._style)
        if not task.no:
            max_no += 1
            task.no = str(max_no)
        _write_task_row(ws, last_row, cols, task, today, bot_comments.get(task.id, []), full=True)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build_workbook(name: str, code: str, event_date: date | None, tasks: list[Task], today: date,
                   bot_comments: dict[int, list[str]] | None = None) -> bytes:
    """ERP-таблица с нуля — для проектов, созданных в боте без загрузки файла.

    Формат совместим с импортом: файл можно доработать в Excel и загрузить обратно.
    """
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ERP"
    header_fill = PatternFill("solid", fgColor="1F3864")
    for col, (_, header, width) in enumerate(STANDARD_HEADERS, start=1):
        cell = ws.cell(1, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    cols = {key: i for i, (key, _, _) in enumerate(STANDARD_HEADERS, start=1)}
    bot_comments = bot_comments or {}
    for i, task in enumerate(tasks, start=1):
        if not task.no:
            task.no = str(i)
        _write_task_row(ws, i + 1, cols, task, today, bot_comments.get(task.id, []), full=True)

    last = max(len(tasks) + 1, 200)
    status_dv = DataValidation(type="list", formula1='"' + ",".join(STATUS_LABELS.values()) + '"', allow_blank=True)
    priority_dv = DataValidation(type="list", formula1='"Критично,Высокий,Средний,Низкий"', allow_blank=True)
    ws.add_data_validation(status_dv)
    ws.add_data_validation(priority_dv)
    status_letter = openpyxl.utils.get_column_letter(cols["status"])
    priority_letter = openpyxl.utils.get_column_letter(cols["priority"])
    status_dv.add(f"{status_letter}2:{status_letter}{last}")
    priority_dv.add(f"{priority_letter}2:{priority_letter}{last}")

    info = wb.create_sheet("Информация о проекте")
    info["A1"] = "Название мероприятия:"
    info["C1"] = name
    info["A2"] = "Код проекта:"
    info["C2"] = code
    info["A3"] = "Дата мероприятия:"
    info["C3"] = event_date.strftime("%d.%m.%Y") if event_date else "Уточнить"
    info.column_dimensions["A"].width = 26
    info.column_dimensions["C"].width = 50

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
